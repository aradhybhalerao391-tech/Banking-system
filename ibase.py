from openpyxl import Workbook, load_workbook
import os

FILE = "bank.xlsx"


def add_new():
    name = input("Enter your name: ")
    phone_no = input("Enter your phone no.: ")
    balance = float(input("Enter your balance: "))

    if os.path.exists(FILE):
        wb = load_workbook(FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Phone No.", "Balance"])

    ws.append([name, phone_no, balance])
    wb.save(FILE)

    print("Account created successfully!")


def deposit():
    num = input("Enter your phone no.: ")
    money = float(input("Enter money to deposit: "))

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[1].value == num:
            row[2].value += money
            wb.save(FILE)
            print("Your money has been successfully deposited!")
            return

    print("Account not found!")


def withdraw():
    num = input("Enter your phone no.: ")
    money = float(input("Enter money to withdraw: "))

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[1].value == num:
            if row[2].value >= money:
                row[2].value -= money
                wb.save(FILE)
                print("Your money has been successfully withdrawn!")
            else:
                print("Insufficient balance!")
            return

    print("Account not found!")


def view_balance():
    num = input("Enter your phone no.: ")

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[1].value == num:
            print("Your balance is:", row[2].value)
            return

    print("Account not found!")


def menu():
    while True:
        print("\nWELCOME TO THE BANK")
        print("1. Add new account")
        print("2. Deposit money")
        print("3. Withdraw money")
        print("4. View balance")
        print("5. Exit")

        ch = int(input("Enter your choice: "))

        if ch == 1:
            add_new()
        elif ch == 2:
            deposit()
        elif ch == 3:
            withdraw()
        elif ch == 4:
            view_balance()
        elif ch == 5:
            print("Thank you...")
            break
        else:
            print("Enter valid choice!")


menu()
